# SPDX-FileCopyrightText: 2026 Keresztes Zsolt <https://kereszteszsolt.hu>
# SPDX-License-Identifier: Apache-2.0

import io
import zipfile

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from pypdf import PdfReader

from contest_check import reports
from contest_check.presentation import policy_rows, summary_rows
from contest_check.presentation import test_rows as detail_rows
from contest_check.results import BatchResult


def pdf_text(payload):
    return "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(payload)).pages)


def test_pdf_has_every_test_units_limits_and_separate_policy_sections(report_batch):
    payload = reports.build_report("pdf", report_batch)
    text = pdf_text(payload)
    for expected in (
        "addition",
        "1001",
        "1002",
        "1003",
        "Per-test outcomes",
        "Wall (s)",
        "CPU (s)",
        "Peak memory",
        "0.5",
        "128",
        "1024",
        "Disqualified / Rejected",
        "Requires Review",
        "N/A",
        "NOT_RUN",
        "source-byte limit exceeded",
        "teacher review required",
    ):
        assert expected in text
    assert text.index("Disqualified / Rejected") < text.index("Requires Review")
    reader = PdfReader(io.BytesIO(payload))
    assert not reader.attachments and "/JavaScript" not in reader.trailer["/Root"]
    assert len(payload) <= reports.MAX_REPORT_BYTES


def test_xlsx_has_all_sheets_exact_text_ids_and_no_active_content(report_batch):
    payload = reports.build_report("xlsx", report_batch)
    workbook = load_workbook(io.BytesIO(payload), data_only=False)
    assert workbook.sheetnames == ["Summary", "Tests", "Disqualified - Rejected", "Requires Review"]
    sheet = workbook["Tests"]
    assert sheet.max_row == 7  # All 3 participants x 2 tests, including not run.
    assert [sheet.cell(row, 2).value for row in range(2, 8)] == ["z", "a"] * 3
    assert sheet["A2"].value == "1001" and sheet["A2"].data_type == "s"
    assert sheet["F2"].value == 0.125 and sheet["G2"].value == "N/A"
    assert sheet["D6"].value == "NOT_RUN"
    for sheet in workbook:
        for row in sheet:
            assert all(cell.data_type != "f" and cell.hyperlink is None for cell in row)
    with zipfile.ZipFile(io.BytesIO(payload)) as archive:
        assert not any(
            "vba" in name.lower() or "externalLink" in name for name in archive.namelist()
        )
    workbook.close()


@pytest.mark.parametrize(
    "value",
    [
        "=1+1",
        "+SUM(1,2)",
        "-2+3",
        "@SUM(1,2)",
        "  =1+1",
        "\t1+1",
        "\r=1+1",
        "\n=1+1",
    ],
)
def test_spreadsheet_text_neutralizes_formula_prefixes(value):
    assert reports.excel_text(value).startswith("'")


def test_spreadsheet_controls_and_safe_identifiers():
    assert reports.excel_text("1000000000000001") == "1000000000000001"
    assert reports.excel_text("plain\x00text\x01") == "plaintext"


@pytest.mark.parametrize("report_format", ["pdf", "xlsx"])
def test_report_bytes_are_deterministic_and_make_no_temporary_files(
    report_batch, report_format, tmp_path, monkeypatch
):
    def forbidden(*args, **kwargs):
        pytest.fail("Report generation must remain in memory")

    monkeypatch.setattr("tempfile.mkstemp", forbidden)
    monkeypatch.setattr("tempfile.NamedTemporaryFile", forbidden)
    first = reports.build_report(report_format, report_batch)
    assert first == reports.build_report(report_format, report_batch)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize("report_format", ["pdf", "xlsx"])
def test_export_limit_does_not_leave_artifacts(report_batch, report_format, monkeypatch, tmp_path):
    monkeypatch.setattr(reports, "MAX_REPORT_BYTES", 100)
    with pytest.raises(reports.ReportLimit):
        reports.build_report(report_format, report_batch)
    assert list(tmp_path.iterdir()) == []


@pytest.mark.parametrize(
    "field, value",
    [
        ("task_id", "<img src='https://invalid.example/secret'>"),
        ("task_id", '=HYPERLINK("https://invalid.example")'),
        ("participants", []),
        ("test_ids", ["a", "z"]),
        ("test_ids", ["z", "z"]),
        ("tests_per_participant", 101),
    ],
)
def test_report_boundary_rejects_inconsistent_or_untrusted_result_fields(
    report_batch, field, value
):
    unsafe = report_batch.model_copy(update={field: value})
    with pytest.raises(reports.ReportError):
        reports.build_report("pdf", unsafe)


def test_result_counts_cannot_be_fabricated(report_batch):
    value = report_batch.model_dump()
    value["participants"][0]["passed_tests"] = 50
    with pytest.raises(ValidationError):
        BatchResult.model_validate(value)


def test_reports_cannot_accept_private_payload_fields(report_batch):
    value = report_batch.model_dump()
    value["source_files"] = {"main.py": "PRIVATE_SOURCE"}
    value["expected_output"] = "PRIVATE_EXPECTED"
    with pytest.raises(ValidationError):
        BatchResult.model_validate(value)


def test_incomplete_batch_lists_all_unrun_tests_and_never_turns_them_into_failures(report_batch):
    value = report_batch.model_dump()
    value["status"] = "incomplete"
    first, *remaining = value["participants"]
    first.update(status="incomplete", passed_tests=1, failed_tests=0, system_errors=1)
    first["tests"][1].update(verdict="SYSTEM_ERROR", reason="sandbox_failure", wall_seconds=None)
    for participant in remaining:
        participant.update(
            status="not_run",
            passed_tests=0,
            failed_tests=0,
            system_errors=0,
            not_run_tests=2,
            tests=[],
            findings=[],
        )
    batch = BatchResult.model_validate(value)
    rows = detail_rows(batch)
    assert len(rows) == 6 and sum(row[2] == "Not run" for row in rows) == 4
    assert summary_rows(batch)[0][2:5] == [1, 0, 1]
    text = pdf_text(reports.build_report("pdf", batch))
    assert "Incomplete batch" in text and "NOT_RUN" in text
    assert policy_rows(batch, "DQ") == []


def test_sixteen_digit_participant_ids_remain_text(report_batch):
    value = report_batch.model_dump()
    value["participants"][0]["participant_id"] = "1000000000000001"
    batch = BatchResult.model_validate(value)
    workbook = load_workbook(io.BytesIO(reports.build_report("xlsx", batch)))
    assert workbook["Tests"]["A2"].value == "1000000000000001"
    assert workbook["Tests"]["A2"].data_type == "s"
    workbook.close()


def test_pdf_paginates_and_repeats_headers(report_batch):
    value = report_batch.model_dump()
    value["tests_per_participant"] = 100
    value["test_ids"] = [f"test_{index}" for index in range(100)]
    participant = value["participants"][0]
    participant.update(passed_tests=100, failed_tests=0)
    participant["tests"] = [
        dict(participant["tests"][0], test_id=test_id) for test_id in value["test_ids"]
    ]
    value["participants"] = [participant]
    payload = reports.build_report("pdf", BatchResult.model_validate(value))
    reader = PdfReader(io.BytesIO(payload))
    assert len(reader.pages) >= 3
    assert "test_99" in pdf_text(payload)
    assert sum("Wall (s)" in page.extract_text() for page in reader.pages) >= 2


@pytest.mark.parametrize("report_format", ["pdf", "xlsx"])
def test_maximum_batch_fits_the_report_contract(report_batch, report_format):
    value = report_batch.model_dump()
    value["tests_per_participant"] = 100
    value["test_ids"] = [f"test_{index}" for index in range(100)]
    participant = value["participants"][0]
    participant.update(passed_tests=100, failed_tests=0)
    participant["tests"] = [
        dict(participant["tests"][0], test_id=test_id) for test_id in value["test_ids"]
    ]
    value["participants"] = [
        dict(participant, participant_id=str(1000 + index)) for index in range(100)
    ]
    payload = reports.build_report(report_format, BatchResult.model_validate(value))
    assert 1000 < len(payload) <= reports.MAX_REPORT_BYTES
    if report_format == "xlsx":
        workbook = load_workbook(io.BytesIO(payload), read_only=True)
        assert workbook["Tests"].max_row == 10_001
        assert workbook["Tests"]["A10001"].value == "1099"
        workbook.close()
    else:
        reader = PdfReader(io.BytesIO(payload))
        text = reader.pages[-1].extract_text()
        assert "1099" in text and "test_99" in text
